#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Malla curricular completa —doce trimestres— de las diez licenciaturas del
expediente de modificación de julio de 2026, normalizada en `mallas.json`
para `construir_mallas.py`.

Cada coordinación entregó su malla en un formato distinto, así que hay un
lector por licenciatura. Todos devuelven lo mismo: trimestre, clave, nombre
tal como viene, créditos y, cuando la malla lo rotula, el tronco.

| Licenciatura | Fuente |
|---|---|
| Ambiental, Química, Física, Eléctrica | `malla_<clave>.json` de los decks de consulta, ya validados al crédito |
| Civil | hoja «Mapa curricular», renglón de claves y créditos, renglón de nombres |
| Industrial, Metalúrgica | hoja de malla con clave, tronco y créditos por celda |
| Mecánica | texto del mapa curricular v35 (`RAG_ModificacionesJulio2026/data/raw`) |
| Computación | Figura 5 de la propuesta de modificaciones, en texto |
| Electrónica | formas del dibujo del boligrama (el `.xlsx` no trae celdas con texto) |

Cada UEA se empareja con el catálogo del tablero (`../datos.js`), del que se
toman el nombre y la ruta de su ficha. Los créditos son los de la malla, que es
lo que suma cada trimestre; cuando el catálogo difiere, se registra.

Uso:  python3 extraer_mallas.py
"""

import difflib
import json
import os
import re
import unicodedata
import zipfile
from collections import Counter, defaultdict

import openpyxl
from lxml import etree

AQUI = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.abspath(os.path.join(AQUI, "..", ".."))
EXP = os.path.join(BASE, "Modificaciones Licenciaturas Julio 2026")
RAW = os.path.join(BASE, "RAG_ModificacionesJulio2026/data/raw")
DECKS = os.path.join(BASE, "Presentaciones_Modificaciones")
SALIDA = os.path.join(AQUI, "mallas.json")

TG = {
    "1114054": "Matemáticas Básicas para Ingeniería",
    "1114048": "Cálculo Diferencial",
    "1114049": "Cálculo Integral",
    "1114056": "Sistemas de Ecuaciones Lineales y Matrices",
    "1114052": "Fundamentos de Física y Aplicaciones I",
    "1114053": "Fundamentos de Física y Aplicaciones II",
    "1114055": "Química General para Ingeniería",
    "1100211": "Comunicación Asertiva",
    "1100212": "Cultura de Paz y Género",
    "1100215": "Vida Universitaria",
}

ETIQUETA = {"TG": "general", "TBP": "profesional", "TP": "profesional",
            "TI": "integracion", "TI OBL": "integracion", "TI OPT": "integracion"}

OPTATIVA = "Optativa del Tronco de Integración"


def sin_marcador(s):
    """«xxxxxxx Seminario de …», «X Estructuras …»: el marcador de clave
    pendiente no es nombre."""
    return re.sub(r"^x+\s+", "", str(s or ""), flags=re.I)


def norm(s):
    s = unicodedata.normalize("NFD", str(s or "").lower())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]+", "", s)


def limpia(s):
    return re.sub(r"\s+", " ", str(s or "")).strip()


def reg(trim, clave=None, nombre=None, creditos=None, tronco=None, opt=False,
        fijo=False):
    clave = limpia(clave) or None
    if clave and not re.fullmatch(r"\d{6,7}|\d{3}[\dX*]{3,4}X?", clave):
        clave = None                     # «TBP», «#######», «NOTA»: sin clave
    return {"trim": int(trim), "clave": clave, "nombre": limpia(nombre) or None,
            "creditos": float(creditos) if creditos not in (None, "") else None,
            "tronco": tronco, "opt": bool(opt), "fijo": fijo}


# ── lectores ─────────────────────────────────────────────────────────────────
def desde_deck(carpeta, clave):
    """Mallas ya normalizadas y cuadradas al crédito para los decks de consulta.
    Sus nombres ya se corrigieron a mano contra el plan, así que se conservan
    (`fijo`); del catálogo se toma sólo la ruta de la ficha."""
    m = json.load(open(os.path.join(DECKS, carpeta, f"malla_{clave}.json"),
                       encoding="utf-8"))
    return [reg(t["trimestre"], nombre=u["nombre"], creditos=u["creditos"],
                tronco=u["tronco"], opt=u.get("opt"), fijo=True)
            for t in m["trimestres"] for u in t["ueas"]]


def malla_civil():
    """Renglón de claves (clave, créditos cada tres columnas), trimestre en la
    columna 19, y en el renglón siguiente los nombres. Seminario y Proyecto de
    Integración llevan «#######» en lugar de clave y las optativas «TBP»."""
    ws = openpyxl.load_workbook(os.path.join(
        EXP, "Ing. Civil/Mapa curricular_claves_Ing_Civil (julio_2026)V4.xlsx"),
        data_only=True)["Mapa curricular"]
    filas = list(ws.iter_rows(values_only=True))
    out = []
    for i, f in enumerate(filas[:-1]):
        t = f[19] if len(f) > 19 else None
        if not isinstance(t, (int, float)) or not 1 <= t <= 12:
            continue
        nombres = filas[i + 1]
        for j in range(1, 19, 3):
            cr, nom = f[j + 1], nombres[j]
            if not isinstance(cr, (int, float)) or not nom:
                continue
            opt = bool(re.match(r"(?i)optativa\s*\d", str(nom)))
            out.append(reg(t, f[j], OPTATIVA if opt else nom, cr,
                           tronco="integracion" if opt else None, opt=opt))
    return out


def _celdas_con_tronco(ruta, hoja, col_trim):
    """Hojas donde cada UEA ocupa tres celdas —clave, tronco, créditos— y el
    renglón siguiente trae el nombre y, en `col_trim`, el trimestre."""
    ws = openpyxl.load_workbook(ruta, data_only=True)[hoja]
    filas = list(ws.iter_rows(values_only=True))
    out = []
    for i, f in enumerate(filas[:-1]):
        nombres = filas[i + 1]
        t = nombres[col_trim] if col_trim < len(nombres) else None
        try:
            t = int(str(t).strip())
        except ValueError:
            continue
        if not 1 <= t <= 12:
            continue
        for j in range(1, len(f) - 1):
            et = limpia(f[j]).upper()
            if et not in ETIQUETA or j + 1 >= 36:
                continue
            try:
                cr = float(str(f[j + 1]).strip())
            except ValueError:
                continue
            nom = nombres[j - 1] if j - 1 < len(nombres) else None
            if not cr or not nom:
                continue
            opt = et == "TI OPT" or "optativa" in norm(nom)
            if opt and norm(nom).startswith("optativasdeltronco"):
                # «Optativas del Tronco de Integración» con 18 créditos: dos casillas
                out += [reg(t, None, OPTATIVA, cr / 2, "integracion", True)] * 2
                continue
            out.append(reg(t, f[j - 1], OPTATIVA if opt else nom, cr,
                           ETIQUETA[et], opt))
    return out


def malla_industrial():
    return _celdas_con_tronco(os.path.join(
        EXP, "Ing. Industrial/5. Boligrama/Diagrama de seriación Industrial  14 09 26.xlsm"),
        "Malla Curr Presenta 7", 36)


def malla_metalurgica():
    return _celdas_con_tronco(os.path.join(
        EXP, "Ing. Metalúrgica/Boligrama_Metalurgia y Materiales_4 sept 26.xlsx"),
        "Mapa Curricular", 1)


def malla_mecanica():
    """Texto del mapa v35: renglón de «clave tronco créditos», y más abajo el
    renglón que empieza con el número de trimestre."""
    f = os.path.join(RAW, "Ing. Mecánica__5. Boligrama__Mapa Curricular_2026 - v35.txt")
    out, pend = [], []
    for linea in open(f, encoding="utf-8", errors="ignore"):
        cl = re.findall(r"(\d{6,7}|\d{3}\w{3}X?|11\*\d{4})\s+(TG|TP|TI OBL|TI OPT)\s+(\d{1,2})\b",
                        linea)
        if len(cl) >= 2:
            pend = cl
            continue
        m = re.match(r"\s*(\d{1,2})\s+\d{1,3}\s", linea)
        if m and pend:
            for c, et, cr in pend:
                opt = et == "TI OPT"
                out.append(reg(m.group(1), None if opt else c, OPTATIVA if opt else None,
                               cr, ETIQUETA[et], opt))
            pend = []
    return out


def malla_computacion():
    """Figura 5 de la propuesta. Las cuatro optativas del Tronco de Integración
    ocupan dos casillas en el trimestre XI y dos en el XII, y la figura exige
    «como mínimo 32 créditos» de ellas; los totales de trimestre (33 y 34) sólo
    cuadran con 8 créditos por casilla."""
    f = os.path.join(RAW, "Ing. Computación__1. Pertinencia y justificación__"
                     "Propuesta de Modificaciones al Plan y Programas de Estudio "
                     "de la Lic. en Ing. en Computación.txt")
    rom = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5, "VI": 6, "VII": 7,
           "VIII": 8, "IX": 9, "X": 10, "XI": 11, "XII": 12}
    texto = open(f, encoding="utf-8", errors="ignore").read()
    texto = texto[:texto.index("Figura 5: Mapa curricular propuesto")]
    bloques = re.split(r"\bTrimestre\s+(XII|XI|X|IX|VIII|VII|VI|V|IV|III|II|I)\b", texto)
    out = []
    for i in range(1, len(bloques) - 1, 2):
        t = rom[bloques[i]]
        for m in re.finditer(r"\b(\d{6,7})\s+(TG|TP|TI)\s+(\d{1,2})\b", bloques[i + 1]):
            out.append(reg(t, m.group(1), None, m.group(3), ETIQUETA[m.group(2)]))
    for t in (11, 11, 12, 12):
        out.append(reg(t, None, OPTATIVA, 8, "integracion", True))
    return out


def malla_electronica():
    """El boligrama de Electrónica es un dibujo: cada UEA es un grupo de formas
    (nombre, clave, créditos) anclado a una celda. Los trimestres ocupan pares
    de renglones —el renglón alto de la etiqueta y el bajo que lo precede—, de
    modo que el trimestre es el renglón de anclaje entre dos."""
    ns = {"x": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
          "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    z = zipfile.ZipFile(os.path.join(EXP, "Ing. Electrónica/05_BoligramaSisEloCiber_Ago26.xlsx"))
    arbol = etree.fromstring(z.read("xl/drawings/drawing1.xml"))
    out = []
    for an in arbol.xpath("/x:wsDr/*", namespaces=ns):
        fr = an.find("x:from", ns)
        if fr is None:
            continue
        r, c = int(fr.find("x:row", ns).text), int(fr.find("x:col", ns).text)
        textos = []
        for sp in an.iter(f"{{{ns['x']}}}sp"):
            x = " ".join("".join(p.xpath(".//a:t/text()", namespaces=ns))
                         for p in sp.xpath(".//a:p", namespaces=ns)).strip()
            if x:
                textos.append(x)
        if not (2 <= r <= 25 and c < 34 and len(textos) == 3):
            continue
        nom, clave, cr = textos
        nom = re.sub(r"\s*\(\d\S*T,\s*[\d.]+P\)\s*$", "", nom)
        if cr == "663":                       # tres cifras superpuestas en la forma
            cr = "6"
        opt = nom.startswith("Optativa")
        out.append(reg(r // 2, clave, nom, cr, "integracion" if opt else None, opt))
    return out


LECTORES = [
    ("amb", lambda: desde_deck("Ambiental", "amb")),
    ("civ", malla_civil),
    ("com", malla_computacion),
    ("ele", lambda: desde_deck("Electrica", "ele")),
    ("elo", malla_electronica),
    ("fis", lambda: desde_deck("Física", "fis")),
    ("ind", malla_industrial),
    ("mec", malla_mecanica),
    ("met", malla_metalurgica),
    ("qui", lambda: desde_deck("Quimica", "qui")),
]

# Claves que la malla escribe mal y que por eso resolverían a otra UEA.
CLAVE_ERRATA = {
    ("elo", "11500034"): "1150044",       # Resolución de Problemas y Pensamiento Crítico
}

# Nombre para las pocas UEA que la malla sólo da por clave provisional.
PROVISIONALES = {
    ("mec", "114002X"): "Estructuras Mecánicas",
    # Computación: la Figura 5 da sólo clave, y estas tres no son únicas en el
    # catálogo o no están en él con esa clave.
    ("com", "1150032"): "Gestión y seguimiento de proyectos tecnológicos",
    ("com", "110046"): "Seminario de integración en ingeniería en computación",
    ("com", "110047"): "Proyecto de integración en ingeniería en computación",
}


# Nombre del catálogo para UEA que la malla titula de otro modo y que el
# emparejamiento por parecido no alcanza.
ALIAS = {
    ("civ", "seminariodeintegracioneningenieriacivil"): "Seminario de Proyecto de Integración",
    ("qui", "seminariodeintegracioneningquimica"): "Seminario de Integración",
    ("qui", "proyectodeintegracioneningenieriaquimica"): "Proyecto de Integración I",
}

# Erratas del propio catálogo que no deben pasar a la malla.
ERRATAS = {"ingeniaría": "ingeniería"}


# ── emparejamiento con el catálogo del tablero ───────────────────────────────
def catalogo():
    t = open(os.path.join(AQUI, "..", "datos.js"), encoding="utf-8").read()
    return json.loads(t[t.index("{"):t.rindex("}") + 1])["licenciaturas"]


def nombre_visible(s):
    """Nombres en mayúsculas sostenidas pasan a caja de oración; los números
    romanos finales conservan su caja."""
    s = limpia(s)
    if s.upper() != s:
        return s[0].upper() + s[1:]
    pal = s.lower().split()
    pal = [p.upper() if re.fullmatch(r"(i|ii|iii|iv|v|vi)", p) and k else p
           for k, p in enumerate(pal)]
    return " ".join(pal)[:1].upper() + " ".join(pal)[1:]


def ruta_ficha(lic, u, ueas):
    """Dirección de la ficha en el tablero. La clave sólo sirve cuando es única
    en la licenciatura: «sin clave», «1130XXX» y las claves tentativas
    repetidas se resuelven por el nombre normalizado, que el tablero entiende
    con el prefijo «~»."""
    iguales = [x for x in ueas if x["clave"] == u["clave"]]
    if len(iguales) == 1 and re.fullmatch(r"\d{6,7}", u["clave"]):
        return f"#/uea/{lic}/{u['clave']}"
    return f"#/uea/{lic}/~{norm(sin_marcador(u['nombre']))}"


def emparejar(lic, r, ueas, por_clave, por_nombre):
    clave = CLAVE_ERRATA.get((lic, r["clave"]), r["clave"])
    # Una UEA del TG se reconoce primero por su nombre: Electrónica rotula
    # Cultura de Paz y Género con la clave de Comunicación Asertiva.
    if r["nombre"]:
        for c, oficial in TG.items():
            if difflib.SequenceMatcher(None, norm(r["nombre"]), norm(oficial)).ratio() > 0.85:
                clave = c
                break
    if clave in TG:
        cand = [u for u in ueas if u["clave"] == clave]
        return cand[0] if cand else None, clave
    if not r["nombre"] and (lic, clave) in PROVISIONALES:
        r = dict(r, nombre=PROVISIONALES[(lic, clave)])
    if clave and len(por_clave.get(clave, [])) == 1:
        u = por_clave[clave][0]
        # la clave se acepta si el nombre no la contradice
        if not r["nombre"] or difflib.SequenceMatcher(
                None, norm(r["nombre"]), norm(u["nombre"])).ratio() > 0.55:
            return u, clave
    n = norm(r["nombre"] or PROVISIONALES.get((lic, clave), ""))
    n = norm(ALIAS.get((lic, n), n)) if (lic, n) in ALIAS else n
    if not n:
        return None, clave
    if n in por_nombre:
        u = por_nombre[n]
        return u, (u["clave"] if u["clave"] in TG else clave)
    cerca = difflib.get_close_matches(n, list(por_nombre), n=1, cutoff=0.8)
    if cerca:
        u = por_nombre[cerca[0]]
        return u, (u["clave"] if u["clave"] in TG else clave)
    # TG por nombre, cuando la malla no trae clave
    for c, oficial in TG.items():
        if difflib.SequenceMatcher(None, n, norm(oficial)).ratio() > 0.85:
            cand = [u for u in ueas if u["clave"] == c]
            return (cand[0] if cand else None), c
    return None, clave


def consolidar():
    cat = {l["clave"]: l for l in catalogo()}
    salida, avisos = {}, []
    for lic, lector in LECTORES:
        l = cat[lic]
        ueas = l["ueas"]
        por_clave = defaultdict(list)
        for u in ueas:
            por_clave[u["clave"]].append(u)
        por_nombre = {norm(sin_marcador(u["nombre"])): u for u in ueas}
        trims = defaultdict(list)
        for r in lector():
            if r["opt"]:
                u, clave = None, None
            else:
                u, clave = emparejar(lic, r, ueas, por_clave, por_nombre)
            if clave in TG:
                nombre, tronco = TG[clave], "general"
            elif r["opt"]:
                nombre, tronco = OPTATIVA, "integracion"
            else:
                nombre = (r["nombre"] if r["fijo"] else None) or \
                    (u["nombre"] if u else None) or r["nombre"] or \
                    PROVISIONALES.get((lic, clave)) or clave
                nombre = sin_marcador(nombre)
                for mal, bien in ERRATAS.items():
                    nombre = nombre.replace(mal, bien)
                tronco = r["tronco"] or (u or {}).get("tronco")
                if tronco not in ("profesional", "integracion", "general"):
                    tronco = "integracion" if re.search(
                        r"(?i)integraci[oó]n|optativa", nombre) else "profesional"
                    avisos.append(f"{lic}: tronco inferido «{tronco}» para {nombre}")
                if tronco == "general":       # el TG son sólo sus diez UEA
                    tronco = "profesional"
            cr = r["creditos"] or (u or {}).get("creditos")
            if u and u.get("creditos") and r["creditos"] and u["creditos"] != r["creditos"] \
                    and clave not in TG:
                avisos.append(f"{lic}: {nombre} — malla {r['creditos']:g} cr, "
                              f"catálogo {u['creditos']:g} cr")
            if not u and not r["opt"]:
                avisos.append(f"{lic}: sin ficha en el tablero — {nombre} ({clave})")
            trims[r["trim"]].append({
                "nombre": nombre_visible(nombre), "clave": clave, "creditos": cr,
                "tronco": tronco, "opt": r["opt"],
                "ruta": ruta_ficha(lic, u, ueas) if u else None,
            })
        salida[lic] = {
            "clave": lic,
            "nombre": l.get("nombre_propuesto") or l["nombre"],
            "creditos_plan": l["creditos"]["propuesto"],
            "trimestres": [{"trimestre": t, "creditos": sum(u["creditos"] or 0 for u in us),
                            "ueas": us} for t, us in sorted(trims.items())],
        }
    return salida, avisos


def validar(mallas):
    print(f"\n{'':4s} {'UEA':>4s} {'TG':>4s} {'TP':>4s} {'TI':>4s} {'suma':>5s} {'plan':>5s}  trimestres")
    for lic, m in mallas.items():
        c = Counter()
        for t in m["trimestres"]:
            for u in t["ueas"]:
                c[u["tronco"]] += u["creditos"] or 0
        n = sum(len(t["ueas"]) for t in m["trimestres"])
        tot, plan = sum(c.values()), m["creditos_plan"]["total"]
        marca = "" if (tot == plan and c["general"] == 57) else "   <-- difiere"
        print(f"{lic:4s} {n:4d} {c['general']:4.0f} {c['profesional']:4.0f} "
              f"{c['integracion']:4.0f} {tot:5.0f} {plan:5d}  "
              f"{len(m['trimestres'])}{marca}")


if __name__ == "__main__":
    mallas, avisos = consolidar()
    for a in avisos:
        print("  ·", a)
    validar(mallas)
    json.dump(mallas, open(SALIDA, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"\nEscrito {os.path.relpath(SALIDA)}")
